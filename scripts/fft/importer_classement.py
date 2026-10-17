#!/usr/bin/env python3
"""Importe le classement padel FFT mensuel dans Supabase (table fft_classement).

La FFT publie chaque mois le classement national hommes et dames. Ce script lit
le fichier publié (PDF, XLSX ou CSV), en extrait rang / nom / prénom / points /
ligue / club, puis remplace d'un bloc le mois concerné dans la base.

    export DATABASE_URL='postgresql://postgres:...@db.<projet>.supabase.co:5432/postgres'
    python3 importer_classement.py classement_padel_homme_2026_10.pdf --dry-run
    python3 importer_classement.py classement_padel_homme_2026_10.pdf classement_padel_femme_2026_10.pdf

Le genre et le mois sont déduits du nom de fichier (« homme »/« femme »/« dame »,
« 2026_10 » ou « 2026-10 ») ; --genre et --mois les forcent.

Dépendances (pip install -r requirements.txt) : psycopg pour l'écriture,
openpyxl pour le XLSX, pymupdf pour le PDF. Aucune n'est nécessaire en --dry-run
sur un CSV.
"""
import argparse
import csv
import io
import os
import re
import sys
import unicodedata

# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------

def normaliser(txt):
    """Minuscules, sans accents ni ponctuation, espaces simples (comme normaliserIdentite côté JS)."""
    txt = unicodedata.normalize('NFD', str(txt or ''))
    txt = ''.join(c for c in txt if unicodedata.category(c) != 'Mn').lower()
    txt = re.sub(r'[^a-z0-9#]+', ' ', txt)
    return txt.strip()


# En-têtes reconnus, sous leur forme normalisée. Le fichier FFT a changé de mise en
# page au fil des années : on accepte les variantes plutôt qu'un ordre de colonnes figé.
ALIAS = {
    'rang':        ['rang', 'classement', 'class', 'clt', 'position', 'place', '#', 'n', 'no'],
    'nom_complet': ['joueur', 'joueuse', 'nom prenom', 'nom et prenom', 'identite', 'nom prenom joueur'],
    'nom':         ['nom', 'nom de famille'],
    'prenom':      ['prenom'],
    'points':      ['points', 'pts', 'total points', 'nb points', 'nombre de points'],
    'nationalite': ['nationalite', 'nat', 'pays'],
    'ligue':       ['ligue'],
    'club':        ['club', 'nom du club'],
    'licence':     ['licence', 'n licence', 'no licence', 'numero de licence', 'num licence', 'licence fft'],
    'assimile':    ['assimile', 'assimilation', 'assimile e'],
    'nb_tournois': ['nb tournois', 'tournois', 'nombre de tournois', 'nb de tournois'],
}
_PAR_ALIAS = {a: champ for champ, alias in ALIAS.items() for a in alias}


def mapper_colonnes(entetes):
    """Renvoie {champ: index} pour une ligne d'en-tête, ou None si ce n'en est pas une.

    Une ligne est un en-tête si elle contient un rang et un nom (complet ou nom seul).
    """
    mapping = {}
    for i, h in enumerate(entetes):
        champ = _PAR_ALIAS.get(normaliser(h))
        if champ and champ not in mapping:
            mapping[champ] = i
    if 'rang' in mapping and ('nom_complet' in mapping or 'nom' in mapping):
        return mapping
    return None


def entier(val):
    """'1 280' -> 1280, '15e' -> 15, '' / 'NC' / '-' -> None."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return int(val)
    chiffres = re.sub(r'[\s  .]', '', str(val))
    m = re.match(r'^(\d+)', chiffres)
    return int(m.group(1)) if m else None


def separer_nom_prenom(nom_complet):
    """'GUY DE CHAMISSO Olivier' -> ('GUY DE CHAMISSO', 'Olivier').

    La FFT écrit le nom en capitales puis le prénom : on prend les mots en capitales
    du début comme nom. Si tout est en capitales, le dernier mot est le prénom.
    """
    mots = str(nom_complet or '').split()
    if not mots:
        return '', ''

    def en_capitales(mot):
        lettres = [c for c in mot if c.isalpha()]
        return bool(lettres) and all(c.isupper() for c in lettres)

    i = 0
    while i < len(mots) and en_capitales(mots[i]):
        i += 1
    if i == len(mots) and len(mots) > 1:
        i = len(mots) - 1
    if i == 0:
        i = 1
    return ' '.join(mots[:i]), ' '.join(mots[i:])


def oui_non(val):
    v = normaliser(val)
    if v in ('oui', 'o', 'x', '1', 'true', 'vrai', 'a', 'assimile'):
        return True
    if v in ('non', 'n', '0', 'false', 'faux'):
        return False
    return None


def ligne_vers_joueur(cellules, mapping):
    """Transforme une ligne brute en dict prêt à insérer, ou None si inexploitable."""
    def lire(champ):
        i = mapping.get(champ)
        if i is None or i >= len(cellules):
            return None
        v = cellules[i]
        return v.strip() if isinstance(v, str) else v

    rang = entier(lire('rang'))
    if not rang:
        return None
    if 'nom_complet' in mapping:
        nom, prenom = separer_nom_prenom(lire('nom_complet'))
    else:
        nom, prenom = str(lire('nom') or '').strip(), str(lire('prenom') or '').strip()
    if not nom:
        return None
    vide = lambda v: (str(v).strip() or None) if v is not None else None
    return {
        'rang': rang,
        'nom': nom,
        'prenom': prenom,
        'points': entier(lire('points')),
        'nationalite': vide(lire('nationalite')),
        'ligue': vide(lire('ligue')),
        'club': vide(lire('club')),
        'licence': vide(lire('licence')),
        'assimile': oui_non(lire('assimile')),
        'nb_tournois': entier(lire('nb_tournois')),
    }


def extraire_joueurs(tableaux):
    """Parcourt des tableaux (listes de lignes) et renvoie les joueurs reconnus.

    L'en-tête est cherché dans chaque tableau ; un tableau sans en-tête (suite d'une
    page PDF) réutilise le dernier en-tête vu. Les lignes d'en-tête répétées sont
    ignorées. Lève ValueError si aucun en-tête n'est jamais trouvé.
    """
    joueurs, mapping = [], None
    for tableau in tableaux:
        for cellules in tableau:
            cellules = ['' if c is None else c for c in cellules]
            entete = mapper_colonnes(cellules)
            if entete:
                mapping = entete
                continue
            if mapping is None:
                continue
            j = ligne_vers_joueur(cellules, mapping)
            if j:
                joueurs.append(j)
    if mapping is None:
        raise ValueError(
            "aucune ligne d'en-tête reconnue (il faut au moins une colonne rang/classement "
            "et une colonne nom ou joueur). En-têtes acceptés : voir ALIAS dans ce script.")
    return joueurs


# ---------------------------------------------------------------------------
# Lecture des fichiers
# ---------------------------------------------------------------------------

def lire_csv(contenu):
    texte = contenu.decode('utf-8-sig', errors='replace') if isinstance(contenu, bytes) else contenu
    try:
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=';,\t')
    except csv.Error:
        dialecte = _PointVirgule
    return [list(csv.reader(io.StringIO(texte), dialecte))]


class _PointVirgule(csv.excel):
    delimiter = ';'


def lire_xlsx(chemin):
    import openpyxl
    classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    return [[list(r) for r in feuille.iter_rows(values_only=True)] for feuille in classeur.worksheets]


def lire_pdf(chemin):
    # pymupdf : bien plus rapide que pdfplumber sur un PDF de plusieurs milliers de pages
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    tableaux = []
    with pymupdf.open(chemin) as doc:
        for n, page in enumerate(doc, 1):
            for t in page.find_tables().tables:
                tableaux.append(t.extract())
            if n % 200 == 0:
                print(f'  … {n}/{doc.page_count} pages', file=sys.stderr)
    return tableaux


def lire_fichier(chemin):
    ext = os.path.splitext(chemin)[1].lower()
    if ext == '.csv':
        with open(chemin, 'rb') as f:
            return extraire_joueurs(lire_csv(f.read()))
    if ext in ('.xlsx', '.xlsm'):
        return extraire_joueurs(lire_xlsx(chemin))
    if ext == '.pdf':
        return extraire_joueurs(lire_pdf(chemin))
    raise ValueError(f'format non pris en charge : {ext} (attendu .pdf, .xlsx ou .csv)')


def detecter_genre_mois(chemin):
    """Déduit ('H'|'F'|None, 'AAAA-MM-01'|None) du nom de fichier."""
    nom = normaliser(os.path.basename(chemin))
    genre = None
    if re.search(r'\b(femme|femmes|dame|dames|feminin)\b', nom):
        genre = 'F'
    elif re.search(r'\b(homme|hommes|messieurs|masculin)\b', nom):
        genre = 'H'
    m = re.search(r'\b(20\d\d) (0[1-9]|1[0-2])\b', nom)
    mois = f'{m.group(1)}-{m.group(2)}-01' if m else None
    return genre, mois


# ---------------------------------------------------------------------------
# Écriture en base
# ---------------------------------------------------------------------------

COLONNES = ['rang', 'nom', 'prenom', 'points', 'nationalite', 'ligue', 'club',
            'licence', 'assimile', 'nb_tournois']


def charger(dsn, joueurs, mois, genre, source):
    """Remplace le classement (mois, genre) en une transaction : jamais de base à moitié chargée."""
    import psycopg
    with psycopg.connect(dsn) as conn, conn.cursor() as cur:
        cur.execute('DELETE FROM fft_classement WHERE mois = %s AND genre = %s', (mois, genre))
        cols = ', '.join(['mois', 'genre'] + COLONNES)
        with cur.copy(f'COPY fft_classement ({cols}) FROM STDIN') as copie:
            for j in joueurs:
                copie.write_row([mois, genre] + [j[c] for c in COLONNES])
        cur.execute(
            """INSERT INTO fft_imports (mois, genre, nb_lignes, source) VALUES (%s, %s, %s, %s)
               ON CONFLICT (mois, genre) DO UPDATE
               SET nb_lignes = EXCLUDED.nb_lignes, source = EXCLUDED.source, importe_le = now()""",
            (mois, genre, len(joueurs), source))


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    p.add_argument('fichiers', nargs='+')
    p.add_argument('--genre', choices=['H', 'F'], help='force le genre (sinon déduit du nom de fichier)')
    p.add_argument('--mois', help='force le mois, format AAAA-MM (sinon déduit du nom de fichier)')
    p.add_argument('--dry-run', action='store_true', help="lit et affiche un aperçu, n'écrit rien")
    args = p.parse_args(argv)

    dsn = os.environ.get('DATABASE_URL')
    if not args.dry_run and not dsn:
        p.error('DATABASE_URL non défini (chaîne de connexion Postgres du projet Supabase Le Rebond)')

    for chemin in args.fichiers:
        genre_auto, mois_auto = detecter_genre_mois(chemin)
        genre = args.genre or genre_auto
        mois = f'{args.mois}-01' if args.mois else mois_auto
        if not genre or not mois:
            p.error(f'{chemin} : genre ou mois introuvable dans le nom de fichier, utilise --genre / --mois')

        print(f'{os.path.basename(chemin)} → genre {genre}, mois {mois[:7]}')
        joueurs = lire_fichier(chemin)
        if not joueurs:
            p.error(f'{chemin} : aucun joueur lu')
        rangs = [j['rang'] for j in joueurs]
        print(f'  {len(joueurs)} joueurs, rangs {min(rangs)} → {max(rangs)}, '
              f'{sum(1 for j in joueurs if j["points"] is not None)} avec points')
        for j in joueurs[:5]:
            print(f"    #{j['rang']:<6} {j['nom']} {j['prenom']}  {j['points'] or ''}  {j['ligue'] or ''}")

        if args.dry_run:
            continue
        charger(dsn, joueurs, mois, genre, os.path.basename(chemin))
        print('  importé.')
    if not args.dry_run:
        print("Pense à cliquer « Actualiser les classements FFT » dans l'annuaire de l'admin.")


if __name__ == '__main__':
    main()
